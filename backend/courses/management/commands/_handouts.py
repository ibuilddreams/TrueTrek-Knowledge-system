"""Renders a `_course_content` handout spec into a real PDF or DOCX file.

The courses that carry hand-authored curriculum pair two written lessons with a
downloadable handout in every module. Those handouts used to be the one shared
demo document every seeded lesson pointed at ("Dummy PDF file"), which made a
finished-looking module open onto nothing. Each one is now generated from the
same file that authored the module, so the worksheet a student downloads is the
worksheet the lesson describes.

A spec is plain data — see `_course_content/__init__.py` for the shape — and is
rendered twice by the same code path, so the PDF and the DOCX of a given handout
always say the same thing:

    {
        "intro": "One paragraph under the title.",
        "sections": [
            {
                "heading": "Section title",
                "paragraphs": ["..."],                       # optional
                "bullets": ["..."],                          # optional
                "fields": ["Label", ...],                    # optional, ruled blanks
                "table": {"columns": [...], "rows": [[...]]},# optional
            },
            ...
        ],
    }

Within a section the four kinds render in that fixed order. `fields` and `table`
are both drawn as bordered tables so a handout reads as something to fill in
rather than something to read, in the PDF viewer and in the DOCX viewer
(`DocumentLessonViewer.jsx` renders .docx through mammoth, which supports
headings, lists and tables but not page geometry).

reportlab is an optional import: a machine without it still seeds, with the PDF
handouts simply absent, the same way a failed download leaves a lesson without a
file. python-docx is already a hard dependency of this project.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    PDF_AVAILABLE = True
except ImportError:  # pragma: no cover - exercised only where reportlab is absent
    PDF_AVAILABLE = False

# Shared ink, so a handout matches the written lessons beside it rather than
# looking like a different product.
INK = "#1F2321"
MUTED = "#6B7280"
RULE = "#D4D4D8"
ACCENT = "#8A6E2F"

# Height of a blank ruled answer row, in points — tall enough to write a line in
# by hand on a printed copy.
FIELD_ROW_HEIGHT = 26


def _pdf_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "HandoutTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            alignment=TA_LEFT,
            textColor=colors.HexColor(INK),
            spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "HandoutSubtitle",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor(MUTED),
            spaceAfter=10,
        ),
        "intro": ParagraphStyle(
            "HandoutIntro",
            parent=base["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor(INK),
            spaceAfter=14,
        ),
        "heading": ParagraphStyle(
            "HandoutHeading",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor(ACCENT),
            spaceBefore=12,
            spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "HandoutBody",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor(INK),
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "HandoutBullet",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor(INK),
        ),
        "cell": ParagraphStyle(
            "HandoutCell",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor(INK),
        ),
        "cell_head": ParagraphStyle(
            "HandoutCellHead",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor(INK),
        ),
    }


def _pdf_fields(fields, styles, width):
    """Label beside a blank box, one row per field."""
    label_width = min(2.6 * inch, width * 0.42)
    data = [[Paragraph(label, styles["cell"]), ""] for label in fields]
    table = Table(
        data,
        colWidths=[label_width, width - label_width],
        rowHeights=[FIELD_ROW_HEIGHT] * len(data),
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(RULE)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#FAF8F4")),
            ]
        )
    )
    return table


def _pdf_table(spec, styles, width):
    columns = spec["columns"]
    header = [Paragraph(text, styles["cell_head"]) for text in columns]
    body = [[Paragraph(str(cell), styles["cell"]) for cell in row] for row in spec["rows"]]
    column_width = width / len(columns)
    table = Table(
        [header] + body,
        colWidths=[column_width] * len(columns),
        rowHeights=[None] + [FIELD_ROW_HEIGHT] * len(body),
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(RULE)),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1EEE7")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 5),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
            ]
        )
    )
    return table


def build_handout_pdf(spec, title, subtitle, destination):
    """Writes the handout to `destination` (a path or a file-like object)."""
    if not PDF_AVAILABLE:  # pragma: no cover - guarded by the caller
        raise RuntimeError("reportlab is not installed")

    styles = _pdf_styles()
    margin = 0.85 * inch
    document = SimpleDocTemplate(
        destination,
        pagesize=LETTER,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=0.8 * inch,
        bottomMargin=0.8 * inch,
        title=title,
        author="Cool Nerds",
        subject=subtitle,
    )
    width = LETTER[0] - 2 * margin

    story = [
        Paragraph(title, styles["title"]),
        Paragraph(subtitle, styles["subtitle"]),
        Paragraph(spec["intro"], styles["intro"]),
    ]

    for section in spec["sections"]:
        story.append(Paragraph(section["heading"], styles["heading"]))
        for paragraph in section.get("paragraphs", ()):
            story.append(Paragraph(paragraph, styles["body"]))
        bullets = section.get("bullets", ())
        if bullets:
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(text, styles["bullet"]), leftIndent=14) for text in bullets],
                    bulletType="bullet",
                    bulletFontSize=9,
                    bulletOffsetY=-1,
                    leftIndent=14,
                    spaceAfter=6,
                )
            )
        if section.get("fields"):
            story.append(_pdf_fields(section["fields"], styles, width))
            story.append(Spacer(1, 8))
        if section.get("table"):
            story.append(_pdf_table(section["table"], styles, width))
            story.append(Spacer(1, 8))

    document.build(story)


def _docx_shade(cell, hex_colour):
    """python-docx has no cell-shading API; this is the documented XML route."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shading)


def build_handout_docx(spec, title, subtitle, destination):
    """Writes the same handout as an editable .docx."""
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Heading 1 rather than the built-in Title style: mammoth maps Title to a
    # plain paragraph, so the DOCX viewer would show the handout's name as body
    # text while every section below it was a real heading.
    document.add_heading(title, level=1)

    sub = document.add_paragraph(subtitle)
    sub_run = sub.runs[0]
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    intro = document.add_paragraph(spec["intro"])
    intro.runs[0].italic = True

    for section in spec["sections"]:
        document.add_heading(section["heading"], level=2)
        for paragraph in section.get("paragraphs", ()):
            document.add_paragraph(paragraph)
        for bullet in section.get("bullets", ()):
            document.add_paragraph(bullet, style="List Bullet")

        if section.get("fields"):
            table = document.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for label in section["fields"]:
                row = table.add_row()
                row.cells[0].text = label
                row.cells[0].paragraphs[0].runs[0].bold = True
                row.cells[0].width = Inches(2.6)
                row.cells[1].width = Inches(3.9)
                _docx_shade(row.cells[0], "FAF8F4")
            document.add_paragraph()

        if section.get("table"):
            columns = section["table"]["columns"]
            table = document.add_table(rows=1, cols=len(columns))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for index, column in enumerate(columns):
                cell = table.rows[0].cells[index]
                cell.text = column
                cell.paragraphs[0].runs[0].bold = True
                _docx_shade(cell, "F1EEE7")
            for values in section["table"]["rows"]:
                row = table.add_row()
                for index, value in enumerate(values):
                    row.cells[index].text = str(value)
            document.add_paragraph()

    document.save(destination)
